# -*- coding: utf-8 -*-
"""Markdown → Word(docx)（WPS 可开）。"""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def md_to_docx(md_path, docx_path):
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"
    st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    def add_table(rows):
        rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        rows = [r for r in rows if not all(set(x) <= set("-: ") for x in r)]
        if not rows:
            return
        t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                cell = t.cell(i, j)
                cell.text = c
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.name = "微软雅黑"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        doc.add_paragraph()

    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.lstrip().startswith("|") and i + 1 < len(lines) and lines[i + 1].lstrip().startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            add_table(tbl)
            continue
        if ln.startswith("### "):
            doc.add_heading(ln[4:], level=3)
        elif ln.startswith("## "):
            doc.add_heading(ln[3:], level=2)
        elif ln.startswith("# "):
            h = doc.add_heading(ln[2:], level=1)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif ln.startswith("- "):
            doc.add_paragraph(ln[2:], style="List Bullet")
        elif ln[:2] in ("1.", "2.", "3.", "4.", "5.", "6."):
            doc.add_paragraph(ln[3:], style="List Number")
        else:
            doc.add_paragraph(ln)
        i += 1
    doc.save(str(docx_path))
